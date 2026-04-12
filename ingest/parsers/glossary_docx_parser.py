"""
AirPlus Assist — Portal Glossary DOCX Parser

Parses the AirPlus Portal Glossary Word document.
Structure: one table with 4 columns:
  Col 0: Letter (ignored)
  Col 1: Topic (= term name, stored as question_text)
  Col 2: Description (= the answer content)
  Col 3: More information (stored as-is in metadata)

One table row = one chunk.
Images are ignored in MVP.

TODO: Phase — production image extraction
"""

from pathlib import Path
from datetime import datetime, timezone
from typing import List, Dict
from docx import Document


def parse_glossary_docx(
    file_path: Path, product: str
) -> List[Dict]:
    """
    Parse the Portal Glossary Word file.
    Returns list of {"content": str, "metadata": dict}
    """
    doc = Document(file_path)
    chunks = []
    ingested_at = datetime.now(timezone.utc).isoformat()
    source_file = file_path.name

    if not doc.tables:
        return chunks

    table = doc.tables[0]
    chunk_index = 0

    for i, row in enumerate(table.rows):
        # Skip header row
        if i == 0:
            continue

        cells = [cell.text.strip() for cell in row.cells]

        # Expect 4 columns
        if len(cells) < 3:
            continue

        # Col 0: Letter — ignored
        topic       = cells[1].strip() if len(cells) > 1 else ""
        description = cells[2].strip() if len(cells) > 2 else ""
        more_info   = cells[3].strip() if len(cells) > 3 else ""

        # Skip empty rows
        if not topic and not description:
            continue

        # Build chunk content
        content_parts = []
        if topic:
            content_parts.append(f"Term: {topic}")
        if description:
            content_parts.append(f"Definition: {description}")
        if more_info:
            content_parts.append(f"Further reading: {more_info}")

        content       = "\n".join(content_parts)
        chunk_preview = content[:120]

        chunks.append({
            "content": content,
            "metadata": {
                "product":          product,
                "source_file":      source_file,
                "source_type":      "glossary_docx",
                "page_number":      None,
                "sheet_name":       None,
                "question_text":    topic,
                "url":              None,
                "section_title":    "Portal Glossary",
                "chunk_preview":    chunk_preview,
                "ingested_at":      ingested_at,
                "chunk_index":      chunk_index,
                "more_information": more_info,
                "search_stage":     "1",
            }
        })
        chunk_index += 1

    return chunks
