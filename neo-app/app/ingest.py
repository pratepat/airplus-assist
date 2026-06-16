"""
Document ingestion: parse all supported formats, chunk, embed via Azure OpenAI,
and upload to Azure AI Search.

Supported formats: PDF, DOCX, XLSX, TXT (including Q&A format), URLs (urls.txt),
                   Glossary DOCX (specialized table parser)
"""
import hashlib
import json
import logging
from datetime import datetime, timezone
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter
from openai import AzureOpenAI
from azure.search.documents import SearchClient
from azure.search.documents.indexes import SearchIndexClient
from azure.search.documents.indexes.models import (
    HnswAlgorithmConfiguration,
    SearchableField,
    SearchField,
    SearchFieldDataType,
    SearchIndex,
    SemanticConfiguration,
    SemanticField,
    SemanticPrioritizedFields,
    SemanticSearch,
    SimpleField,
    VectorSearch,
    VectorSearchProfile,
)
from azure.core.credentials import AzureKeyCredential

from .config import config

log = logging.getLogger(__name__)

# ── Azure clients (module-level, created lazily) ───────────────────────────────

_openai_client: AzureOpenAI | None = None
_search_client: SearchClient | None = None
_index_client: SearchIndexClient | None = None


def _get_openai() -> AzureOpenAI:
    global _openai_client
    if _openai_client is None:
        _openai_client = AzureOpenAI(
            azure_endpoint=config.azure_openai_endpoint,
            api_key=config.azure_openai_api_key or None,
            api_version=config.azure_openai_api_version,
        )
    return _openai_client


def _get_search_client() -> SearchClient:
    global _search_client
    if _search_client is None:
        _search_client = SearchClient(
            endpoint=config.azure_search_endpoint,
            index_name=config.azure_search_index,
            credential=AzureKeyCredential(config.azure_search_key),
        )
    return _search_client


def _get_index_client() -> SearchIndexClient:
    global _index_client
    if _index_client is None:
        _index_client = SearchIndexClient(
            endpoint=config.azure_search_endpoint,
            credential=AzureKeyCredential(config.azure_search_key),
        )
    return _index_client


# ── Index management ───────────────────────────────────────────────────────────

def ensure_index() -> None:
    """Create the Azure AI Search index if it doesn't already exist."""
    client = _get_index_client()
    existing = set(client.list_index_names())
    if config.azure_search_index in existing:
        log.info("Index '%s' already exists", config.azure_search_index)
        return

    log.info("Creating index '%s' …", config.azure_search_index)
    index = SearchIndex(
        name=config.azure_search_index,
        fields=[
            SimpleField(name="id", type=SearchFieldDataType.String, key=True),
            SearchableField(name="content", type=SearchFieldDataType.String),
            SimpleField(name="product",       type=SearchFieldDataType.String, filterable=True, retrievable=True),
            SimpleField(name="source_file",   type=SearchFieldDataType.String, retrievable=True),
            SimpleField(name="source_type",   type=SearchFieldDataType.String, filterable=True, retrievable=True),
            SimpleField(name="search_stage",  type=SearchFieldDataType.String, filterable=True, retrievable=True),
            SimpleField(name="page_number",   type=SearchFieldDataType.Int32,  retrievable=True),
            SimpleField(name="sheet_name",    type=SearchFieldDataType.String, retrievable=True),
            SearchableField(name="question_text", type=SearchFieldDataType.String, retrievable=True),
            SimpleField(name="url",            type=SearchFieldDataType.String, retrievable=True),
            SimpleField(name="section_title",  type=SearchFieldDataType.String, retrievable=True),
            SimpleField(name="more_information", type=SearchFieldDataType.String, retrievable=True),
            SimpleField(name="ingested_at",    type=SearchFieldDataType.String, retrievable=True),
            SimpleField(name="chunk_index",    type=SearchFieldDataType.Int32,  retrievable=True),
            SearchField(
                name="content_vector",
                type=SearchFieldDataType.Collection(SearchFieldDataType.Single),
                searchable=True,
                vector_search_dimensions=1536,
                vector_search_profile_name="hnsw-profile",
            ),
        ],
        vector_search=VectorSearch(
            algorithms=[HnswAlgorithmConfiguration(name="hnsw")],
            profiles=[VectorSearchProfile(name="hnsw-profile", algorithm_configuration_name="hnsw")],
        ),
        semantic_search=SemanticSearch(
            configurations=[
                SemanticConfiguration(
                    name="semantic_config",
                    prioritized_fields=SemanticPrioritizedFields(
                        content_fields=[SemanticField(field_name="content")],
                        title_field=SemanticField(field_name="question_text"),
                    ),
                )
            ]
        ),
    )
    client.create_index(index)
    log.info("Index '%s' created", config.azure_search_index)


# ── Embedding ──────────────────────────────────────────────────────────────────

def embed_texts(texts: list[str]) -> list[list[float]]:
    """Embed a list of texts using Azure OpenAI. Batches up to 100 per call."""
    client = _get_openai()
    vectors: list[list[float]] = []
    batch_size = 100
    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        response = client.embeddings.create(
            input=batch,
            model=config.azure_openai_embed_model,
        )
        vectors.extend(item.embedding for item in response.data)
        log.info("  Embedded %d/%d", min(i + batch_size, len(texts)), len(texts))
    return vectors


# ── Upload to Azure AI Search ──────────────────────────────────────────────────

def upload_chunks(chunks: list[dict], vectors: list[list[float]]) -> None:
    """Merge chunks + vectors and upload to Azure AI Search in batches of 100."""
    client = _get_search_client()
    documents = []
    for chunk, vector in zip(chunks, vectors):
        meta = chunk["metadata"]
        doc_id = hashlib.md5(
            f"{meta.get('product')}:{meta.get('source_file')}:{meta.get('chunk_index')}:{chunk['content'][:50]}".encode()
        ).hexdigest()
        documents.append({
            "id":               doc_id,
            "content":          chunk["content"],
            "product":          meta.get("product") or "",
            "source_file":      meta.get("source_file") or "",
            "source_type":      meta.get("source_type") or "",
            "search_stage":     meta.get("search_stage") or "",
            "page_number":      meta.get("page_number"),
            "sheet_name":       meta.get("sheet_name") or "",
            "question_text":    meta.get("question_text") or "",
            "url":              meta.get("url") or "",
            "section_title":    meta.get("section_title") or "",
            "more_information": meta.get("more_information") or "",
            "ingested_at":      meta.get("ingested_at") or "",
            "chunk_index":      meta.get("chunk_index") or 0,
            "content_vector":   vector,
        })

    batch_size = 100
    for i in range(0, len(documents), batch_size):
        batch = documents[i : i + batch_size]
        result = client.upload_documents(documents=batch)
        failed = [r for r in result if not r.succeeded]
        if failed:
            log.warning("  %d document(s) failed to upload in batch %d", len(failed), i // batch_size)
    log.info("Uploaded %d documents to index '%s'", len(documents), config.azure_search_index)


# ── Helpers ────────────────────────────────────────────────────────────────────

def _splitter() -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=config.chunk_size,
        chunk_overlap=config.chunk_overlap,
    )


def _base_meta(product: str, source_file: str, source_type: str) -> dict:
    return {
        "product":          product,
        "source_file":      source_file,
        "source_type":      source_type,
        "page_number":      None,
        "sheet_name":       None,
        "question_text":    None,
        "url":              None,
        "section_title":    None,
        "more_information": None,
        "search_stage":     None,
        "ingested_at":      datetime.now(timezone.utc).isoformat(),
        "chunk_index":      0,
    }


def _load_stage_config(product_dir: Path) -> dict | None:
    """Load stage_config.json from a product directory, or return None if absent."""
    cfg_path = product_dir / "stage_config.json"
    if not cfg_path.exists():
        return None
    try:
        with open(cfg_path) as fh:
            return json.load(fh)
    except Exception as exc:
        log.warning("Failed to read %s: %s", cfg_path, exc)
        return None


# ── Parsers (unchanged from local version) ────────────────────────────────────

def parse_pdf(file_path: Path, product: str) -> list[dict]:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        log.warning("PyMuPDF not installed — skipping %s", file_path.name)
        return []

    chunks = []
    sp = _splitter()
    try:
        doc = fitz.open(str(file_path))
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text()
            if not text.strip():
                continue
            for i, piece in enumerate(sp.split_text(text)):
                meta = _base_meta(product, file_path.name, "pdf")
                meta["page_number"] = page_num
                meta["chunk_index"] = i
                chunks.append({"content": piece, "metadata": meta})
        doc.close()
    except Exception as e:
        log.error("PDF parse error %s: %s", file_path.name, e)
    return chunks


def parse_docx(file_path: Path, product: str) -> list[dict]:
    try:
        from docx import Document
    except ImportError:
        log.warning("python-docx not installed — skipping %s", file_path.name)
        return []

    chunks = []
    sp = _splitter()
    try:
        doc = Document(str(file_path))
        current_section: str | None = None
        buffer: list[str] = []

        def flush():
            if not buffer:
                return
            text = "\n".join(buffer)
            for i, piece in enumerate(sp.split_text(text)):
                meta = _base_meta(product, file_path.name, "docx")
                meta["section_title"] = current_section
                meta["chunk_index"] = i
                chunks.append({"content": piece, "metadata": meta})

        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            if style.startswith("Heading"):
                flush()
                buffer.clear()
                current_section = text
            else:
                buffer.append(text)
        flush()
    except Exception as e:
        log.error("DOCX parse error %s: %s", file_path.name, e)
    return chunks


def parse_glossary_docx(file_path: Path, product: str) -> list[dict]:
    """
    Parse a portal-style glossary Word document.
    Expects a table with 4 columns: Letter | Topic | Description | More Information.
    One row → one chunk, tagged as search_stage="1".
    """
    try:
        from docx import Document
    except ImportError:
        log.warning("python-docx not installed — skipping %s", file_path.name)
        return []

    chunks = []
    ingested_at = datetime.now(timezone.utc).isoformat()
    try:
        doc = Document(str(file_path))
        if not doc.tables:
            return chunks
        table = doc.tables[0]
        for i, row in enumerate(table.rows):
            if i == 0:
                continue
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) < 3:
                continue
            topic       = cells[1] if len(cells) > 1 else ""
            description = cells[2] if len(cells) > 2 else ""
            more_info   = cells[3] if len(cells) > 3 else ""
            if not topic and not description:
                continue
            content_parts = []
            if topic:
                content_parts.append(f"Term: {topic}")
            if description:
                content_parts.append(f"Definition: {description}")
            if more_info:
                content_parts.append(f"Further reading: {more_info}")
            content = "\n".join(content_parts)
            meta = _base_meta(product, file_path.name, "glossary_docx")
            meta["question_text"]    = topic
            meta["section_title"]    = "Glossary"
            meta["more_information"] = more_info or None
            meta["search_stage"]     = "1"
            meta["ingested_at"]      = ingested_at
            meta["chunk_index"]      = i
            chunks.append({"content": content, "metadata": meta})
    except Exception as e:
        log.error("Glossary DOCX parse error %s: %s", file_path.name, e)
    return chunks


def parse_excel(file_path: Path, product: str) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        log.warning("openpyxl not installed — skipping %s", file_path.name)
        return []

    chunks = []
    try:
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            key_idx = next((i for i, h in enumerate(headers) if h.lower() == "key"), None)
            if key_idx is None:
                continue

            content_cols = [
                (i, h) for i, h in enumerate(headers)
                if i != key_idx
                and not h.lower().rstrip("?").endswith("approved")
            ]

            for row_num, row in enumerate(rows[1:], start=2):
                key_val = row[key_idx] if key_idx < len(row) else None
                if not key_val:
                    continue
                key_val = str(key_val).strip()

                parts = [f"Term: {key_val}"]
                for i, col_name in content_cols:
                    val = row[i] if i < len(row) else None
                    if val is not None and str(val).strip():
                        parts.append(f"{col_name}: {str(val).strip()}")

                meta = _base_meta(product, file_path.name, "xlsx")
                meta["sheet_name"]   = sheet_name
                meta["question_text"] = key_val
                meta["chunk_index"]  = row_num
                chunks.append({"content": "\n".join(parts), "metadata": meta})
        wb.close()
    except Exception as e:
        log.error("XLSX parse error %s: %s", file_path.name, e)
    return chunks


def parse_txt(file_path: Path, product: str) -> list[dict]:
    try:
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1")
    except Exception as e:
        log.error("TXT read error %s: %s", file_path.name, e)
        return []

    is_qa = text.count("\nQ:") > 3 or text.startswith("Q:")
    chunks = []

    if is_qa:
        current_q: str | None = None
        current_a: list[str] = []
        pairs: list[tuple[str, str]] = []

        for line in text.splitlines():
            if line.startswith("Q:"):
                if current_q is not None:
                    pairs.append((current_q, "\n".join(current_a)))
                current_q = line[2:].strip()
                current_a = []
            elif line.startswith("A:"):
                current_a.append(line[2:].strip())
            else:
                current_a.append(line)
        if current_q:
            pairs.append((current_q, "\n".join(current_a)))

        for i, (q, a) in enumerate(pairs):
            meta = _base_meta(product, file_path.name, "txt")
            meta["question_text"] = q
            meta["section_title"] = "FAQ"
            meta["chunk_index"]   = i
            chunks.append({"content": f"Q: {q}\nA: {a}", "metadata": meta})
    else:
        sp = _splitter()
        for i, piece in enumerate(sp.split_text(text)):
            meta = _base_meta(product, file_path.name, "txt")
            meta["chunk_index"] = i
            chunks.append({"content": piece, "metadata": meta})

    return chunks


def parse_urls(file_path: Path, product: str) -> list[dict]:
    try:
        import trafilatura
    except ImportError:
        log.warning("trafilatura not installed — skipping %s", file_path.name)
        return []

    try:
        raw = file_path.read_text(encoding="utf-8")
        urls = [u.strip() for u in raw.splitlines() if u.strip() and not u.startswith("#")]
    except Exception as e:
        log.error("URL file read error %s: %s", file_path.name, e)
        return []

    chunks = []
    sp = _splitter()

    for url in urls:
        try:
            downloaded = trafilatura.fetch_url(url)
            if not downloaded:
                continue
            text = trafilatura.extract(downloaded)
            if not text or not text.strip():
                continue
            meta_info = trafilatura.extract_metadata(downloaded)
            title = meta_info.title if meta_info and meta_info.title else url
            for i, piece in enumerate(sp.split_text(text)):
                meta = _base_meta(product, file_path.name, "url")
                meta["url"]           = url
                meta["section_title"] = title
                meta["chunk_index"]   = i
                chunks.append({"content": piece, "metadata": meta})
            log.info("  URL %s → %d chunks", url, i + 1)
        except Exception as e:
            log.warning("URL fetch error %s: %s", url, e)

    return chunks


# ── Ingest summary (returned instead of IngestResult) ─────────────────────────

class IngestSummary:
    """Lightweight result returned after ingest — no embeddings stored locally."""
    products: list[str]
    total: int
    ingested_at: str

    def __init__(self):
        self.products    = []
        self.total       = 0
        self.ingested_at = ""


# ── Ingest orchestrator ────────────────────────────────────────────────────────

def ingest_documents(docs_path: str | Path) -> IngestSummary:
    """
    Scan docs_path/{product}/ folders, parse all supported documents,
    embed via Azure OpenAI, and upload to Azure AI Search.

    The index is created if it doesn't exist. Existing documents are
    overwritten (upsert by deterministic id).
    """
    docs_path = Path(docs_path)
    summary   = IngestSummary()

    if not docs_path.exists():
        log.warning("Documents path does not exist: %s", docs_path)
        return summary

    ensure_index()

    all_chunks: list[dict] = []
    products: set[str] = set()

    for product_dir in sorted(docs_path.iterdir()):
        if not product_dir.is_dir() or product_dir.name.startswith("."):
            continue

        product = product_dir.name
        products.add(product)
        product_chunks: list[dict] = []

        stage_cfg    = _load_stage_config(product_dir)
        stage1_files = set(stage_cfg.get("stage1", [])) if stage_cfg else set()

        for file_path in sorted(product_dir.iterdir()):
            if file_path.name.startswith(".") or not file_path.is_file():
                continue
            if file_path.suffix.lower() == ".json":
                continue

            ext  = file_path.suffix.lower()
            name = file_path.name.lower()

            if ext == ".docx" and name.startswith("glossary"):
                fc = parse_glossary_docx(file_path, product)
            elif ext == ".pdf":
                fc = parse_pdf(file_path, product)
            elif ext == ".docx":
                fc = parse_docx(file_path, product)
            elif ext in (".xlsx", ".xls"):
                fc = parse_excel(file_path, product)
            elif ext == ".txt" and name == "urls.txt":
                fc = parse_urls(file_path, product)
            elif ext == ".txt":
                fc = parse_txt(file_path, product)
            else:
                log.debug("Skipping unsupported file: %s", file_path.name)
                continue

            if stage_cfg:
                for chunk in fc:
                    if chunk["metadata"].get("search_stage") is None:
                        chunk["metadata"]["search_stage"] = (
                            "1" if file_path.name in stage1_files else "2"
                        )

            log.info("  %s → %d chunks", file_path.name, len(fc))
            product_chunks.extend(fc)

        log.info("Product '%s': %d chunks total", product, len(product_chunks))
        all_chunks.extend(product_chunks)

    if not all_chunks:
        log.warning("No chunks found in %s", docs_path)
        return summary

    log.info("Embedding %d chunks via Azure OpenAI …", len(all_chunks))
    texts   = [c["content"] for c in all_chunks]
    vectors = embed_texts(texts)

    upload_chunks(all_chunks, vectors)

    summary.products    = sorted(products)
    summary.total       = len(all_chunks)
    summary.ingested_at = datetime.now(timezone.utc).isoformat()

    log.info(
        "Ingest complete: %d chunks from %d product(s): %s",
        summary.total, len(summary.products), summary.products,
    )
    return summary
